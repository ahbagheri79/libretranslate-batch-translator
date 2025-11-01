import re
import time
import json
import math
import os
from pathlib import Path
from typing import List, Tuple
import pandas as pd
import requests

prop_path = Path("/home/bagheri/Desktop/Untitled-1.properties")
out_path = Path("/home/bagheri/Desktop/translations.xlsx")

LT_URL = os.getenv("LT_URL", "http://[::1]:5000/translate")
LT_API_KEY = os.getenv("LT_API_KEY", "")
SRC_LANG = "en"
DST_LANG = "fa"

MIN_INTERVAL_SEC = 0.1
DEFAULT_RETRY_AFTER = 150
ESCAPE_ASCII_ONLY = False


def to_java_unicode_escapes(s: str) -> str:
    if not s:
        return ""
    be = s.encode("utf-16-be")
    return "".join("\\u%04X" % int.from_bytes(be[i:i + 2], "big") for i in range(0, len(be), 2))


def to_ascii_with_unicode_escapes(s: str) -> str:
    if not s:
        return ""
    out = []
    for ch in s:
        if ord(ch) < 128:
            out.append(ch)
        else:
            be = ch.encode("utf-16-be")
            for i in range(0, len(be), 2):
                out.append("\\u%04X" % int.from_bytes(be[i:i + 2], "big"))
    return "".join(out)


PLACEHOLDER_PATTERNS = [
    r"\{[^\}]+\}",
    r"%[sdif]",
    r"\$\{[^\}]+\}",
]


def _mask_placeholders(text: str) -> Tuple[str, List[str]]:
    if not text:
        return text, []
    patterns = [re.compile(p) for p in PLACEHOLDER_PATTERNS]
    originals = []

    def repl(m):
        originals.append(m.group(0))
        return f"__PH_{len(originals) - 1}__"

    masked = text
    for pat in patterns:
        masked = pat.sub(repl, masked)
    return masked, originals


def _unmask_placeholders(text: str, originals: List[str]) -> str:
    if not text or not originals:
        return text
    for i, ph in enumerate(originals):
        text = text.replace(f"__PH_{i}__", ph)
    return text


def translate_text(text: str, session: requests.Session) -> str:
    if not text:
        return ""
    masked, originals = _mask_placeholders(text)
    payload = {
        "q": masked,
        "source": SRC_LANG,
        "target": DST_LANG,
        "format": "text",
    }
    if LT_API_KEY:
        payload["api_key"] = LT_API_KEY

    resp = session.post(LT_URL, data=payload, timeout=30)
    if resp.status_code == 429:
        retry_after = resp.headers.get("Retry-After")
        raise RateLimitedError(retry_after=retry_after)
    resp.raise_for_status()

    data = resp.json()
    translated = data.get("translatedText", "")
    return _unmask_placeholders(translated, originals)


class RateLimitedError(Exception):
    def __init__(self, retry_after=None):
        self.retry_after = retry_after
        super().__init__("HTTP 429: Rate limited")


def print_progress(done: int, total: int, start_ts: float):
    pct = (done / total) * 100 if total else 100.0
    elapsed = time.time() - start_ts
    speed = done / elapsed if elapsed > 0 else 0.0
    remaining = (total - done) / speed if speed > 0 else 0.0

    def fmt_sec(sec: float) -> str:
        sec = int(sec)
        m, s = divmod(sec, 60)
        h, m = divmod(m, 60)
        if h > 0:
            return f"{h:02d}:{m:02d}:{s:02d}"
        return f"{m:02d}:{s:02d}"

    print(f"\rTranslating: {done}/{total} ({pct:5.1f}%)  "
          f"Elapsed: {fmt_sec(elapsed)}  "
          f"Remaining: {fmt_sec(remaining)}", end="", flush=True)


def main():
    keys, english = [], []
    with prop_path.open("r", encoding="utf-8", errors="ignore") as f:
        for raw in f:
            line = raw.strip()
            if not line or line.startswith("#"):
                continue
            m = re.match(r"([^=:#\s][^=:]*?)\s*[:=]\s*(.*)", line)
            if not m:
                continue
            k, v = m.group(1).strip(), m.group(2).strip()
            v = v.replace("''", "'")
            keys.append(k)
            english.append(v)

    df = pd.DataFrame({"key": keys, "english": english, "fa": ["" for _ in keys]})

    total = len(df)
    if total == 0:
        print("No rows found.")
        return

    session = requests.Session()
    start_ts = time.time()
    last_req_ts = 0.0
    translated_fa = [""] * total

    print(f"Starting translation of {total} rows...\n(Request interval: {MIN_INTERVAL_SEC}s)")
    print_progress(0, total, start_ts)

    i = 0
    while i < total:
        src_text = df.at[i, "fa"] or df.at[i, "english"]

        if not src_text:
            translated_fa[i] = ""
            i += 1
            print_progress(i, total, start_ts)
            continue

        delta = time.time() - last_req_ts
        if delta < MIN_INTERVAL_SEC:
            time.sleep(MIN_INTERVAL_SEC - delta)

        try:
            translated = translate_text(df.at[i, "english"] if not df.at[i, "fa"] else df.at[i, "fa"], session)
            translated_fa[i] = translated
            last_req_ts = time.time()
            i += 1
            print_progress(i, total, start_ts)

        except RateLimitedError as e:
            retry_after = e.retry_after
            if retry_after is not None:
                try:
                    wait_sec = int(retry_after)
                except ValueError:
                    wait_sec = DEFAULT_RETRY_AFTER
            else:
                wait_sec = DEFAULT_RETRY_AFTER

            print("\n⚠️  Rate limit applied (HTTP 429).")
            print(f"⏳ Waiting {wait_sec} seconds before resuming...")
            for s in range(wait_sec, 0, -1):
                print(f"\rResuming in: {s:4d}s", end="", flush=True)
                time.sleep(1)
            print("\n▶️ Resuming...")
            last_req_ts = time.time()

        except requests.RequestException as e:
            print(f"\n❗️ Network/HTTP error at row {i + 1}: {e}. Original text kept.")
            translated_fa[i] = src_text
            last_req_ts = time.time()
            i += 1
            print_progress(i, total, start_ts)

    print("\nTranslation complete. Building UTF-16 column...")

    df["fa"] = translated_fa
    source_for_utf = df["fa"].where(df["fa"].astype(str).str.len() > 0, df["english"])
    converter = to_ascii_with_unicode_escapes if ESCAPE_ASCII_ONLY else to_java_unicode_escapes
    df["utf16"] = source_for_utf.apply(converter)

    engine = None
    try:
        import xlsxwriter
        engine = "xlsxwriter"
    except Exception:
        try:
            import openpyxl
            engine = "openpyxl"
        except Exception:
            engine = None

    if engine:
        with pd.ExcelWriter(out_path, engine=engine) as writer:
            df.to_excel(writer, index=False, sheet_name="translations")
    else:
        out_path_csv = out_path.with_suffix(".csv")
        df.to_csv(out_path_csv, index=False)
        print(f"⚠️ Excel dependencies not installed; CSV created: {out_path_csv}")
        return

    print(f"✅ File created: {out_path}")


if __name__ == "__main__":
    main()
